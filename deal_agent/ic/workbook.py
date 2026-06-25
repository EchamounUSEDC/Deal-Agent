"""Build the live IC Go/No-Go Excel dashboard.

Self-contained — no API key, no add-in. Two ways to add a deal and get a verdict:

1. **Type it in.** The Deals sheet ships with ready "New Deal" slots. Pick the deal type
   and the market (a county dropdown), enter ~10 economics, and the GO / CONDITIONAL /
   NO-GO verdict computes instantly against the Springfield & Hamburg statics. The market
   stats (population, growth, pipeline, rank) auto-fill from the county via VLOOKUP.
2. **Import a pro forma** with `python -m deal_agent.ic.cli add --proforma X.xlsx --market ...`,
   or the one-click VBA button in the macro-enabled build (see excel/AddDeal.bas).

Every gate formula reproduces `schema.py` exactly. Sheets: Dashboard · Read Me · Add a Deal ·
Deals · Scores · Ranking · Market Ranking · Market Data.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule

from .schema import Deal, INCOME, DEVELOPMENT
from .market import COL, MarketRanking

# ---- styling ---------------------------------------------------------------------
NAVY = "1F3A5F"
WHITE = "FFFFFF"
GREEN = "C6EFCE"; GREEN_T = "006100"
AMBER = "FFEB9C"; AMBER_T = "9C6500"
RED = "FFC7CE"; RED_T = "9C0006"
GREY = "F2F2F2"
INPUT = "FFF2CC"      # yellow — user types here
BENCH = "DDEBF7"      # light blue — locked benchmark
H = Font(bold=True, color=WHITE, size=11)
TITLE = Font(bold=True, color=WHITE, size=16)
B = Font(bold=True)
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEADFILL = PatternFill("solid", fgColor=NAVY)
GREYFILL = PatternFill("solid", fgColor=GREY)
INPUTFILL = PatternFill("solid", fgColor=INPUT)
BENCHFILL = PatternFill("solid", fgColor=BENCH)
GREENFILL = PatternFill("solid", fgColor=GREEN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

# (field, label, kind) — order defines the Deals sheet layout.
METRICS = [
    ("deal_type", "Deal type", "text"),
    ("market_county", "Market (county)", "text"),
    ("is_development", "Is development (1/0)", "int"),
    ("total_cost", "Total cost", "money"),
    ("equity", "Equity", "money"),
    ("debt", "Debt", "money"),
    ("ltv", "LTV", "pct"),
    ("going_in_cap", "Going-in cap", "pct"),
    ("exit_cap", "Exit cap", "pct"),
    ("noi", "NOI", "money"),
    ("opex_ratio", "Op-ex ratio", "pct"),
    ("yield_on_cost", "Yield on cost", "pct"),
    ("dev_spread", "Dev. spread", "pct"),
    ("unlevered_irr", "Unlevered IRR", "pct"),
    ("levered_irr", "Levered IRR", "pct"),
    ("nrsf", "NRSF", "num"),
    ("site_acreage", "Site acreage", "num"),
    ("units", "Units", "num"),
    ("avg_rate_per_nrsf", "Avg rate $/NRSF", "num2"),
    ("market_population", "Market population", "num"),
    ("market_pop_growth_5yr", "5-yr pop growth", "pct"),
    ("market_pipeline_pct", "Supply pipeline %", "pct"),
    ("market_rank", "Market rank", "int"),
    ("market_score", "Market weighted score", "num2"),
    ("best_irr", "Best IRR (calc)", "pct"),
    ("underwriting_yield", "Underwriting yield (calc)", "pct"),
]
MROW = {f: i + 2 for i, (f, _, _) in enumerate(METRICS)}  # field -> Deals row
FMT = {"money": '$#,##0', "pct": '0.0%', "num": '#,##0', "num2": '#,##0.00', "int": '0'}

# Cells a user types for a new deal (everything else is a formula/lookup).
INPUT_FIELDS = {"total_cost", "equity", "debt", "going_in_cap", "exit_cap", "noi",
                "opex_ratio", "unlevered_irr", "levered_irr", "nrsf", "site_acreage",
                "units", "avg_rate_per_nrsf"}

# Scores sheet rows. Mirrors schema.GATES weights/thresholds.
SCORE_ROWS = [
    ("IRR beats exit cap", 2), ("Yield clears floor", 3), ("Op-ex in range", 4),
    ("Population ≥ 200k", 5), ("Positive pop growth", 6), ("Pipeline contained", 7),
    ("Return clears target", 8), ("Market upper half", 9), ("Size near 75k NRSF", 10),
]
S_TOTAL, S_CRIT, S_VERDICT = 12, 13, 14
GATE_WEIGHTS = {2: 15, 3: 15, 4: 10, 5: 12, 6: 10, 7: 10, 8: 13, 9: 8, 10: 7}


def _gate_formula(row: int, dc: str) -> str:
    """Scores-sheet formula for gate `row`, candidate in Deals column `dc`."""
    d = lambda r: f"Deals!{dc}{r}"
    irr, cap = d(MROW["best_irr"]), d(MROW["exit_cap"])
    yld, opex, pop = d(MROW["underwriting_yield"]), d(MROW["opex_ratio"]), d(MROW["market_population"])
    grw, pipe = d(MROW["market_pop_growth_5yr"]), d(MROW["market_pipeline_pct"])
    isdev, rank, nrsf = d(MROW["is_development"]), d(MROW["market_rank"]), d(MROW["nrsf"])
    return {
        2: f'=IF(OR({irr}="",{cap}=""),0,IF({irr}>{cap},15,0))',
        3: f'=IF({yld}="",0,IF({yld}>=0.075,15,IF({yld}>=0.065,7.5,0)))',
        4: f'=IF({opex}="",0,IF({opex}<=0.3,10,IF({opex}<=0.35,5,0)))',
        5: f'=IF({pop}="",0,IF({pop}>=200000,12,0))',
        6: f'=IF({grw}="",5,IF({grw}>=0.05,10,IF({grw}>0,5,0)))',
        7: f'=IF({pipe}="",5,IF({pipe}<=0.05,10,IF({pipe}<=0.1,5,0)))',
        8: f'=IF({irr}="",0,IF({irr}>=IF({isdev}=1,0.15,0.1),13,IF({irr}>=IF({isdev}=1,0.12,0.07),6.5,0)))',
        9: f'=IF({rank}="",4,IF({rank}<=123,8,IF({rank}<=184,4,0)))',
        10: f'=IF({nrsf}="",3.5,IF(AND({nrsf}>=40000,{nrsf}<=120000),7,IF(AND({nrsf}>=25000,{nrsf}<=150000),3.5,0)))',
    }[row]


def _critical_formula(dc: str) -> str:
    d = lambda r: f"Deals!{dc}{r}"
    irr, cap = d(MROW["best_irr"]), d(MROW["exit_cap"])
    yld, pop = d(MROW["underwriting_yield"]), d(MROW["market_population"])
    irr_fail = f'IF(OR({irr}="",{cap}=""),1,IF({irr}>{cap},0,1))'
    yld_fail = f'IF({yld}="",1,IF({yld}>=0.065,0,1))'
    pop_fail = f'IF({pop}="",1,IF({pop}>=200000,0,1))'
    return f'=IF(OR({irr_fail}=1,{yld_fail}=1,{pop_fail}=1),1,0)'


def _set(ws, coord, value, *, font=None, fill=None, fmt=None, align=None, border=True):
    c = ws[coord]
    c.value = value
    if font: c.font = font
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    if align: c.alignment = align
    if border: c.border = BORDER
    return c


def build_workbook(out_path: str, benchmarks: list[Deal], candidates: list[Deal],
                   market: MarketRanking | None = None, template_slots: int = 8) -> str:
    fixed = list(benchmarks) + list(candidates)
    n_fixed = len(fixed)
    n_slots = template_slots
    n = n_fixed + n_slots
    last_letter = get_column_letter(1 + n)
    cols = [get_column_letter(2 + i) for i in range(n)]
    slot_cols = cols[n_fixed:]                       # the editable "New Deal" columns
    n_bench = len(benchmarks)

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True

    # ============================ Market Data (for VLOOKUP) =====================
    n_markets = _build_market_data(wb.create_sheet("Market Data"), market)

    # ============================ Deals (inputs) ================================
    dz = wb.active
    dz.title = "Deals"
    _set(dz, "A1", "Metric", font=H, fill=HEADFILL, align=LEFT)
    for field, label, _ in METRICS:
        _set(dz, f"A{MROW[field]}", label, font=B, fill=GREYFILL, align=LEFT)

    for j, deal in enumerate(fixed):
        _write_fixed_column(dz, cols[j], deal, locked=deal.is_benchmark)
    for k, c in enumerate(slot_cols, 1):
        _write_template_column(dz, c, k, n_markets)

    dz.column_dimensions["A"].width = 24
    for c in cols:
        dz.column_dimensions[c].width = 18
    dz.freeze_panes = "B2"

    # ============================ Scores (live engine) ==========================
    sc = wb.create_sheet("Scores")
    _set(sc, "A1", "Gate", font=H, fill=HEADFILL, align=LEFT)
    for label, r in SCORE_ROWS:
        _set(sc, f"A{r}", f"{label} (max {GATE_WEIGHTS[r]})", font=B, fill=GREYFILL, align=LEFT)
    _set(sc, f"A{S_TOTAL}", "TOTAL SCORE", font=B, fill=GREYFILL, align=LEFT)
    _set(sc, f"A{S_CRIT}", "Critical fail (1=veto)", font=B, fill=GREYFILL, align=LEFT)
    _set(sc, f"A{S_VERDICT}", "VERDICT", font=B, fill=GREYFILL, align=LEFT)
    for c in cols:
        _set(sc, f"{c}1", f"=Deals!{c}1", font=H, fill=HEADFILL, align=CENTER)
        for _, r in SCORE_ROWS:
            _set(sc, f"{c}{r}", _gate_formula(r, c), align=CENTER)
        _set(sc, f"{c}{S_TOTAL}", f"=SUM({c}2:{c}10)", font=B, align=CENTER)
        _set(sc, f"{c}{S_CRIT}", _critical_formula(c), align=CENTER)
        _set(sc, f"{c}{S_VERDICT}",
             f'=IF({c}{S_CRIT}=1,"NO-GO",IF({c}{S_TOTAL}>=70,"GO",'
             f'IF({c}{S_TOTAL}>=50,"CONDITIONAL GO","NO-GO")))', font=B, align=CENTER)
    sc.column_dimensions["A"].width = 26
    for c in cols:
        sc.column_dimensions[c].width = 16
    _verdict_cf(sc, f"{cols[0]}{S_VERDICT}:{cols[-1]}{S_VERDICT}")

    # ============================ named ranges ==================================
    wb.defined_names.add(DefinedName("DealNames", attr_text=f"Deals!$B$1:${last_letter}$1"))
    if n_markets:
        wb.defined_names.add(DefinedName(
            "MarketCounties", attr_text=f"'Market Data'!$A$2:$A${n_markets+1}"))
        wb.defined_names.add(DefinedName(
            "MarketTable", attr_text=f"'Market Data'!$A$2:$F${n_markets+1}"))

    # ============================ Dashboard / Add / Ranking / Market view =======
    db = wb.create_sheet("Dashboard")
    _build_dashboard(db, fixed, cols, last_letter, n_bench)
    _build_add_sheet(wb.create_sheet("Add a Deal"), slot_cols, bool(n_markets))
    _build_ranking(wb.create_sheet("Ranking"), fixed, slot_cols, cols)
    if market is not None:
        _build_market_view(wb.create_sheet("Market Ranking"), market)
    _build_readme(wb.create_sheet("Read Me"))

    # tab order: Dashboard, Read Me, Add a Deal, Deals, Scores, Ranking, ...
    order = ["Dashboard", "Read Me", "Add a Deal", "Deals", "Scores", "Ranking",
             "Market Ranking", "Market Data"]
    wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)
    wb.active = wb.sheetnames.index("Dashboard")
    wb["Market Data"].sheet_state = "hidden"

    wb.save(out_path)
    return out_path


def _write_fixed_column(dz, c, deal: Deal, locked: bool):
    fill = BENCHFILL if locked else None
    _set(dz, f"{c}1", deal.name, font=H, fill=HEADFILL, align=CENTER)
    for field, label, kind in METRICS:
        r = MROW[field]
        if field == "is_development":
            val = 1 if deal.is_development else 0
        elif field == "market_county":
            val = deal.market_label
        elif field == "best_irr":
            _set(dz, f"{c}{r}", f'=IF(COUNT({c}{MROW["unlevered_irr"]}:{c}{MROW["levered_irr"]})=0,"",'
                 f'MAX({c}{MROW["unlevered_irr"]}:{c}{MROW["levered_irr"]}))', fmt=FMT["pct"], align=CENTER)
            continue
        elif field == "underwriting_yield":
            _set(dz, f"{c}{r}", f'=IF({c}{MROW["is_development"]}=1,{c}{MROW["yield_on_cost"]},{c}{MROW["going_in_cap"]})',
                 fmt=FMT["pct"], align=CENTER)
            continue
        else:
            val = getattr(deal, field, None)
        _set(dz, f"{c}{r}", val, fmt=FMT.get(kind), align=CENTER,
             fill=fill if field not in ("best_irr", "underwriting_yield") else None)


def _write_template_column(dz, c, k, n_markets):
    """A ready-to-fill 'New Deal' slot: pick type + county, type economics, verdict computes."""
    cc = lambda f: f"{c}{MROW[f]}"
    _set(dz, f"{c}1", f"New Deal {k}", font=H, fill=HEADFILL, align=CENTER)
    # deal type (dropdown) + county (dropdown)
    _set(dz, cc("deal_type"), INCOME, fill=INPUTFILL, align=CENTER)
    _set(dz, cc("market_county"), "", fill=INPUTFILL, align=CENTER)
    # derived: is_development from the deal-type dropdown
    _set(dz, cc("is_development"), f'=IF(EXACT({cc("deal_type")},"{DEVELOPMENT}"),1,0)', align=CENTER, fmt="0")
    # plain inputs
    for field in INPUT_FIELDS:
        _set(dz, cc(field), None, fill=INPUTFILL, align=CENTER, fmt=FMT.get(dict((f, k2) for f, _, k2 in METRICS)[field]))
    # derived economics
    _set(dz, cc("ltv"), f'=IFERROR({cc("debt")}/{cc("total_cost")},"")', fmt=FMT["pct"], align=CENTER)
    _set(dz, cc("yield_on_cost"), f'=IFERROR({cc("noi")}/{cc("total_cost")},"")', fmt=FMT["pct"], align=CENTER)
    # yield_on_cost is a derived field, so drop the yellow input we may have set
    dz[cc("yield_on_cost")].fill = PatternFill()
    _set(dz, cc("dev_spread"),
         f'=IF(AND({cc("is_development")}=1,{cc("yield_on_cost")}<>"",{cc("exit_cap")}<>""),'
         f'{cc("yield_on_cost")}-{cc("exit_cap")},"")', fmt=FMT["pct"], align=CENTER)
    _set(dz, cc("best_irr"),
         f'=IF(COUNT({cc("unlevered_irr")}:{cc("levered_irr")})=0,"",MAX({cc("unlevered_irr")}:{cc("levered_irr")}))',
         fmt=FMT["pct"], align=CENTER)
    _set(dz, cc("underwriting_yield"),
         f'=IF({cc("is_development")}=1,{cc("yield_on_cost")},{cc("going_in_cap")})', fmt=FMT["pct"], align=CENTER)
    # market stats auto-fill from the county via VLOOKUP
    if n_markets:
        county = cc("market_county")
        for field, col_idx, fmt in [("market_population", 2, FMT["num"]),
                                    ("market_pop_growth_5yr", 3, FMT["pct"]),
                                    ("market_pipeline_pct", 4, FMT["pct"]),
                                    ("market_rank", 5, FMT["int"]),
                                    ("market_score", 6, FMT["num2"])]:
            _set(dz, cc(field), f'=IFERROR(VLOOKUP({county},MarketTable,{col_idx},FALSE),"")',
                 fmt=fmt, align=CENTER)
    # data validations: deal type + county
    dv_type = DataValidation(type="list", formula1=f'"{INCOME},{DEVELOPMENT}"', allow_blank=False)
    dz.add_data_validation(dv_type); dv_type.add(dz[cc("deal_type")])
    if n_markets:
        dv_cty = DataValidation(type="list", formula1="=MarketCounties", allow_blank=True)
        dv_cty.prompt = "Pick the deal's county"; dz.add_data_validation(dv_cty)
        dv_cty.add(dz[cc("market_county")])


def _build_market_data(ms, market: MarketRanking | None) -> int:
    """Flat county lookup table: A=County, B=Pop, C=Growth5yr, D=Pipeline, E=Rank, F=Score."""
    _set(ms, "A1", "County", font=B); _set(ms, "B1", "Population", font=B)
    _set(ms, "C1", "Growth5yr", font=B); _set(ms, "D1", "Pipeline", font=B)
    _set(ms, "E1", "Rank", font=B); _set(ms, "F1", "Score", font=B)
    if market is None:
        return 0
    rows = sorted(market.rows, key=lambda r: float(r.get(COL["rank"], 1e9) or 1e9))
    for i, row in enumerate(rows):
        r = 2 + i
        ms[f"A{r}"] = row.get(COL["county"], "")
        for col, key in zip("BCDEF", ["population", "growth_5yr", "pipeline", "rank", "score"]):
            try:
                ms[f"{col}{r}"] = float(row.get(COL[key], ""))
            except (TypeError, ValueError):
                ms[f"{col}{r}"] = row.get(COL[key], "")
    ms.column_dimensions["A"].width = 34
    return len(rows)


def _verdict_cf(ws, rng: str):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"GO"'],
        fill=PatternFill("solid", fgColor=GREEN), font=Font(bold=True, color=GREEN_T)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"CONDITIONAL GO"'],
        fill=PatternFill("solid", fgColor=AMBER), font=Font(bold=True, color=AMBER_T)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"NO-GO"'],
        fill=PatternFill("solid", fgColor=RED), font=Font(bold=True, color=RED_T)))


def _build_dashboard(db, fixed, cols, last_letter, n_bench):
    bench_cols = cols[:2] if len(cols) >= 2 else cols
    db.sheet_view.showGridLines = False
    db.column_dimensions["A"].width = 26
    for c in ("B", "C", "D", "E"):
        db.column_dimensions[c].width = 19

    db.merge_cells("A1:E1")
    _set(db, "A1", "INVESTMENT COMMITTEE  ·  GO / NO-GO DASHBOARD", font=TITLE,
         fill=HEADFILL, align=CENTER, border=False)
    db.row_dimensions[1].height = 30

    _set(db, "A3", "Select a deal ▼", font=B, fill=GREYFILL, align=LEFT)
    sel = _set(db, "B3", fixed[-1].name, font=Font(bold=True, size=12), align=CENTER)
    sel.fill = INPUTFILL
    _set(db, "H1", "=MATCH($B$3,Deals!$1:$1,0)", align=CENTER, border=False)
    db["H1"].font = Font(color="FFFFFF"); db.column_dimensions["H"].hidden = True

    drange = f"Deals!$A$1:${last_letter}${max(MROW.values())}"
    srange = f"Scores!$A$1:${last_letter}${S_VERDICT}"
    isel = "$H$1"

    _set(db, "A5", "VERDICT", font=B, fill=GREYFILL, align=CENTER)
    _set(db, "B5", f"=INDEX({srange},{S_VERDICT},{isel})", font=Font(bold=True, size=14), align=CENTER)
    _set(db, "A6", "Score / 100", font=B, fill=GREYFILL, align=CENTER)
    _set(db, "B6", f"=INDEX({srange},{S_TOTAL},{isel})", font=Font(bold=True, size=14), align=CENTER)
    _set(db, "C5", "Closest archetype", font=B, fill=GREYFILL, align=CENTER)
    _set(db, "C6", f'=IF(INDEX({drange},{MROW["is_development"]},{isel})=1,'
                   f'"Hamburg (Development)","Springfield (Income/Core)")', font=B, align=CENTER)
    _verdict_cf(db, "B5")

    r0 = 8
    _set(db, f"A{r0}", "Metric", font=H, fill=HEADFILL, align=LEFT)
    _set(db, f"B{r0}", "Selected deal", font=H, fill=HEADFILL, align=CENTER)
    for k, bc in enumerate(bench_cols):
        _set(db, f"{get_column_letter(3+k)}{r0}", f"=Deals!{bc}1", font=H, fill=HEADFILL, align=CENTER)

    show = ["deal_type", "market_county", "total_cost", "equity", "debt", "ltv", "going_in_cap",
            "exit_cap", "noi", "opex_ratio", "yield_on_cost", "dev_spread", "unlevered_irr",
            "levered_irr", "nrsf", "market_population", "market_pop_growth_5yr",
            "market_pipeline_pct", "market_rank", "best_irr"]
    kind_of = {f: k for f, _, k in METRICS}
    label_of = {f: l for f, l, _ in METRICS}
    for i, field in enumerate(show):
        r = r0 + 1 + i
        mr = MROW[field]; fmt = FMT.get(kind_of[field])
        _set(db, f"A{r}", label_of[field], font=B, fill=GREYFILL, align=LEFT)
        _set(db, f"B{r}", f"=INDEX({drange},{mr},{isel})", fmt=fmt, align=CENTER)
        for k, bc in enumerate(bench_cols):
            _set(db, f"{get_column_letter(3+k)}{r}", f"=Deals!{bc}{mr}", fmt=fmt, align=CENTER)

    g0 = r0 + len(show) + 3
    _set(db, f"A{g0}", "Gate", font=H, fill=HEADFILL, align=LEFT)
    _set(db, f"B{g0}", "Points (selected)", font=H, fill=HEADFILL, align=CENTER)
    _set(db, f"C{g0}", "Max", font=H, fill=HEADFILL, align=CENTER)
    for i, (label, sr) in enumerate(SCORE_ROWS):
        r = g0 + 1 + i
        _set(db, f"A{r}", label, font=B, fill=GREYFILL, align=LEFT)
        _set(db, f"B{r}", f"=INDEX({srange},{sr},{isel})", align=CENTER)
        _set(db, f"C{r}", GATE_WEIGHTS[sr], align=CENTER)
        db.conditional_formatting.add(f"B{r}", CellIsRule(operator="greaterThanOrEqual",
            formula=[f"C{r}"], fill=PatternFill("solid", fgColor=GREEN)))
        db.conditional_formatting.add(f"B{r}", CellIsRule(operator="equal",
            formula=["0"], fill=PatternFill("solid", fgColor=RED)))

    # On-sheet copy of the deal names so the dropdown works in every app (Excel
    # desktop/web, Google Sheets, LibreOffice) — a cross-sheet list source does not.
    n_deals = len(cols)
    for i, dcol in enumerate(cols):
        _set(db, f"J{i+1}", f"=Deals!{dcol}1", border=False)
        db[f"J{i+1}"].font = Font(color="FFFFFF")
    db.column_dimensions["J"].hidden = True
    dv = DataValidation(type="list", formula1=f"$J$1:$J${n_deals}", allow_blank=False)
    dv.prompt = "Choose a deal to evaluate"; db.add_data_validation(dv); dv.add(db["B3"])


# 'Add a Deal' form: (label, field, kind, dropdown) at rows FORM_ROW0.. in column B.
FORM_ROW0 = 21
FORM_FIELDS = [
    ("Deal name", "name", "text", None),
    ("Deal type", "deal_type", "text", "type"),
    ("Market (county)", "market_county", "text", "county"),
    ("Total cost", "total_cost", "money", None),
    ("Equity", "equity", "money", None),
    ("Debt", "debt", "money", None),
    ("Going-in cap (income)", "going_in_cap", "pct", None),
    ("Exit cap", "exit_cap", "pct", None),
    ("NOI (stabilized)", "noi", "money", None),
    ("Op-ex ratio", "opex_ratio", "pct", None),
    ("Unlevered IRR", "unlevered_irr", "pct", None),
    ("Levered IRR", "levered_irr", "pct", None),
    ("NRSF", "nrsf", "num", None),
    ("Site acreage", "site_acreage", "num", None),
    ("Units", "units", "num", None),
]


def _build_add_sheet(ad, slot_cols, has_market):
    ad.sheet_view.showGridLines = False
    ad.column_dimensions["A"].width = 50
    ad.column_dimensions["B"].width = 30
    cty = " and pick its county from the dropdown" if has_market else ""
    lines = [
        ("ADD A DEAL", TITLE, HEADFILL),
        ("", None, None),
        ("★ DROP A SPREADSHEET IN — click the 'Import Pro Forma' button (macro build). Pick a", B, GREENFILL),
        ("   pro forma file; it reads the metrics, fills the next slot, and shows the verdict.", B, GREENFILL),
        ("   (Setup once: save as .xlsm, import excel/AddDeal.bas, assign ImportProForma to a button.)", None, None),
        ("", None, None),
        ("Other ways — all decide GO / NO-GO live off the Springfield & Hamburg statics:", B, None),
        ("", None, None),
        ("A) Double-click launchers/'Import Deal' (or drag a file onto it): a file picker opens,", None, None),
        ("   it auto-detects the market and scores the deal. No Excel macros, no command line.", None, None),
        ("B) Type directly on the Deals tab: find the first empty 'New Deal' column (yellow),", None, None),
        (f"   pick Deal type{cty}, and enter the economics. The market stats auto-fill.", None, None),
        ("C) Fill the form below, then click the Add Deal button — copies it into the next slot.", None, None),
        ("", None, None),
        (f"Ready-to-fill slots: {', '.join(slot_cols)} (Deals tab).", Font(italic=True), GREYFILL),
    ]
    for i, (txt, font, fill) in enumerate(lines):
        _set(ad, f"A{i+1}", txt, font=font, fill=fill, align=LEFT, border=False)
        if fill is HEADFILL:
            ad.row_dimensions[i+1].height = 26

    # the input form (read by the AddDeal macro; also fine to copy manually)
    _set(ad, f"A{FORM_ROW0-1}", "NEW DEAL FORM", font=H, fill=HEADFILL, align=LEFT)
    _set(ad, f"B{FORM_ROW0-1}", "enter values →", font=H, fill=HEADFILL, align=CENTER)
    for i, (label, field, kind, dd) in enumerate(FORM_FIELDS):
        r = FORM_ROW0 + i
        _set(ad, f"A{r}", label, font=B, fill=GREYFILL, align=LEFT)
        cell = _set(ad, f"B{r}", INCOME if field == "deal_type" else None,
                    fill=INPUTFILL, align=CENTER, fmt=FMT.get(kind))
        if dd == "type":
            dv = DataValidation(type="list", formula1=f'"{INCOME},{DEVELOPMENT}"', allow_blank=False)
            ad.add_data_validation(dv); dv.add(cell)
        elif dd == "county" and has_market:
            dv = DataValidation(type="list", formula1="=MarketCounties", allow_blank=True)
            ad.add_data_validation(dv); dv.add(cell)


def _build_ranking(rk, fixed, slot_cols, cols):
    _set(rk, "A1", "Go/No-Go Ranking — filled deals score live; sort by Score to rank",
         font=B, border=False)
    hdr = 3
    for j, t in enumerate(["Deal", "Type", "Score", "Verdict"]):
        _set(rk, f"{get_column_letter(1+j)}{hdr}", t, font=H, fill=HEADFILL, align=CENTER)
    types = [d.deal_type for d in fixed] + ["(set on Deals tab)"] * len(slot_cols)
    for i, c in enumerate(cols):
        r = hdr + 1 + i
        _set(rk, f"A{r}", f"=Deals!{c}1", align=LEFT)
        _set(rk, f"B{r}", (f'=Deals!{c}{MROW["deal_type"]}' if i >= len(fixed) else types[i]), align=LEFT)
        _set(rk, f"C{r}", f"=Scores!{c}{S_TOTAL}", align=CENTER)
        _set(rk, f"D{r}", f"=Scores!{c}{S_VERDICT}", font=B, align=CENTER)
    _verdict_cf(rk, f"D{hdr+1}:D{hdr+len(cols)}")
    rk.column_dimensions["A"].width = 22
    rk.column_dimensions["B"].width = 30
    for c in ("C", "D"):
        rk.column_dimensions[c].width = 16


def _build_market_view(ms, market: MarketRanking):
    headers = ["Rank", "State", "CBSA", "County / City", "Population", "5-yr Growth",
               "Pipeline %", "Op-ex", "Exit Cap", "Yield/Cost", "Dev Spread", "IRR", "Score"]
    keys = ["rank", "state", "cbsa", "county", "population", "growth_5yr", "pipeline",
            "opex", "exit_cap", "yield_on_cost", "dev_spread", "irr", "score"]
    fmts = {"Population": '#,##0', "5-yr Growth": '0.0%', "Pipeline %": '0.0%', "Op-ex": '0.0%',
            "Exit Cap": '0.0%', "Yield/Cost": '0.0%', "Dev Spread": '0.0%', "IRR": '0.0%', "Score": '0.000'}
    for j, hname in enumerate(headers):
        _set(ms, f"{get_column_letter(1+j)}1", hname, font=H, fill=HEADFILL, align=CENTER)
    for i, row in enumerate(market.top(60)):
        r = 2 + i
        for j, key in enumerate(keys):
            raw = row.get(COL[key], "")
            try:
                val = float(raw)
            except (TypeError, ValueError):
                val = raw
            _set(ms, f"{get_column_letter(1+j)}{r}", val,
                 fmt=fmts.get(headers[j]), align=CENTER if key != "county" else LEFT)
    for j, w in enumerate([6, 14, 30, 26, 12, 11, 11, 9, 9, 11, 11, 9, 9]):
        ms.column_dimensions[get_column_letter(1+j)].width = w
    ms.freeze_panes = "A2"


def _build_readme(rd):
    rd.sheet_view.showGridLines = False
    rd.column_dimensions["A"].width = 112
    lines = [
        ("INVESTMENT COMMITTEE — GO / NO-GO DASHBOARD", TITLE, HEADFILL),
        ("", None, None),
        ("How to use", B, GREYFILL),
        ("• Dashboard tab: pick a deal from the dropdown (B3). Verdict, comparison vs Springfield", None, None),
        ("  & Hamburg, and gate scores recalculate instantly.", None, None),
        ("• Add a Deal tab: how to drop in a new deal (yellow 'New Deal' columns on the Deals tab).", None, None),
        ("", None, None),
        ("The two proven winners (the benchmark statics)", B, GREYFILL),
        ("• Springfield (Income/Core DST): all-equity, ~7% going-in & exit cap, ~14% op-ex, ~11% IRR.", None, None),
        ("  Durable in-place yield, no leverage or lease-up risk — the conservative core benchmark.", None, None),
        ("• Hamburg (Development): ~40% LTV, ~7.4% yield-on-cost vs ~5.5% exit cap (~190 bps spread),", None, None),
        ("  ~17% IRR. Value created at a yield well above the exit cap — the value-add benchmark.", None, None),
        ("", None, None),
        ("The gates", B, GREYFILL),
        ("Critical (any fail ⇒ NO-GO): IRR > exit cap · yield ≥ 6.5% · population ≥ 200k.", None, None),
        ("Scored: op-ex ≤ 35% · positive 5-yr growth (pref ≥5%) · pipeline ≤ 10% (pref ≤5%) ·", None, None),
        ("IRR clears archetype target (Income ≥10% / Dev ≥15%) · market in upper half · ~75k NRSF.", None, None),
        ("Score ≥ 70 = GO · 50–69 = CONDITIONAL GO · <50 or critical fail = NO-GO.", None, None),
        ("", None, None),
        ("Not investment advice. Verify all figures independently.", Font(italic=True), None),
    ]
    for i, (txt, font, fill) in enumerate(lines):
        _set(rd, f"A{i+1}", txt, font=font, fill=fill, align=LEFT, border=False)
        if fill is HEADFILL:
            rd.row_dimensions[i+1].height = 26
