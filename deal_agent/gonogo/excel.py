"""Live Excel functions for Go/No-Go — ask a cell a question, get a live answer.

Once wired into Excel via the **xlwings add-in**, these become worksheet
functions that recompute on every edit:

    =GONOGO_ASK("is Hamburg a go?")   -> "Hamburg: GO (100/100). ..."
    =GONOGO("TX")                     -> "GO"
    =GONOGO_SCORE("Springfield")      -> 16
    =GONOGO_WHY("NE")                 -> "NO-GO — fails the IRR floor ..."
    =GONOGO_BENCHMARK("Hamburg","yield_on_cost") -> 0.074

Setup (Excel on Windows/Mac):
    pip install "xlwings>=0.30"
    xlwings addin install
    # In the workbook: xlwings ribbon -> set the UDF Modules to this module, then
    # "Import Functions". Or run `build_live_workbook()` below to get a ready
    # template, and `python -m deal_agent.gonogo serve` notes in docs/GO_NO_GO.md.

Each call rebuilds the deal library from the configured files + benchmarks, so
answers are always live. Point it at your deals with the GONOGO_DEAL_FILES env
var or a ./deals folder (see library.py).
"""

from __future__ import annotations

from .benchmarks import BENCHMARKS
from .library import answer, load_library, resolve

try:  # xlwings is optional — the module still imports (and the CLI works) without it
    import xlwings as xw

    _HAS_XW = True
except Exception:  # pragma: no cover
    xw = None
    _HAS_XW = False


def _decision(query: str) -> str:
    v = resolve(query)
    return v.decision if v else "NOT FOUND"


def _score(query: str):
    v = resolve(query)
    return round(v.score, 1) if v else "NOT FOUND"


def _why(query: str) -> str:
    v = resolve(query)
    return v.rationale if v else f"No deal matches {query!r}."


def _benchmark(name: str, metric: str = "yield_on_cost"):
    b = BENCHMARKS.get(name) or BENCHMARKS.get(name.title())
    if not b:
        return f"Unknown benchmark {name!r} (have: {', '.join(BENCHMARKS)})."
    return getattr(b, metric, f"Unknown metric {metric!r}.")


# --- Plain Python entry points (work everywhere) ----------------------------
def gonogo(query: str) -> str:
    """GO / CONDITIONAL / NO-GO for a deal name, question, or file path."""
    return _decision(query)


def gonogo_score(query: str):
    """0-100 Go/No-Go score for a deal name, question, or file path."""
    return _score(query)


def gonogo_why(query: str) -> str:
    """The rationale behind a deal's verdict."""
    return _why(query)


def gonogo_ask(question: str) -> str:
    """Answer a plain-English question, e.g. 'is the Hamburg deal a go?'."""
    return answer(question)


def gonogo_benchmark(name: str, metric: str = "yield_on_cost"):
    """Look up a Hamburg/Lewiston benchmark value (e.g. dev_spread_bps)."""
    return _benchmark(name, metric)


# --- Register the same functions as live xlwings UDFs ------------------------
if _HAS_XW:  # pragma: no cover - exercised only inside Excel
    GONOGO_ASK = xw.func(gonogo_ask)
    GONOGO = xw.func(gonogo)
    GONOGO_SCORE = xw.func(gonogo_score)
    GONOGO_WHY = xw.func(gonogo_why)
    GONOGO_BENCHMARK = xw.func(gonogo_benchmark)


def build_live_workbook(out_path: str = "GoNoGo_Live.xlsx") -> str:
    """Write a ready-to-use 'Ask' workbook pre-wired with the live formulas.

    The workbook works as a static snapshot immediately; to make the cells
    recompute live, install the xlwings add-in and import this module's
    functions (see the Setup tab and docs/GO_NO_GO.md).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    lib = load_library()
    wb = Workbook()

    # --- Ask tab ------------------------------------------------------------
    ws = wb.active
    ws.title = "Ask"
    hdr = PatternFill("solid", fgColor="1F4E78")
    hf = Font(color="FFFFFF", bold=True)
    ws["A1"] = "Ask a Go/No-Go question"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Your question:"
    ws["A3"].font = Font(bold=True)
    ws["B3"] = "is Hamburg a go or no go?"
    ws["B3"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["A4"] = "Live answer:"
    ws["A4"].font = Font(bold=True)
    ws["B4"] = "=GONOGO_ASK(B3)"     # live once xlwings UDFs are imported
    ws["A6"] = "Quick lookups (type a deal name in column A):"
    ws["A6"].font = Font(bold=True)
    for i, (lbl, formula) in enumerate([
        ("Decision", "=GONOGO(A{r})"),
        ("Score", "=GONOGO_SCORE(A{r})"),
        ("Why", "=GONOGO_WHY(A{r})"),
    ]):
        ws.cell(7, 2 + i, lbl).font = hf
        ws.cell(7, 2 + i).fill = hdr
    ws.cell(7, 1, "Deal").font = hf
    ws.cell(7, 1).fill = hdr
    sample = ["Hamburg", "Lewiston", "TX", "NE", "Springfield"]
    for k, name in enumerate(sample):
        r = 8 + k
        ws.cell(r, 1, name)
        ws.cell(r, 2, f"=GONOGO(A{r})")
        ws.cell(r, 3, f"=GONOGO_SCORE(A{r})")
        ws.cell(r, 4, f"=GONOGO_WHY(A{r})")
    ws.column_dimensions["A"].width = 22
    for col in ("B", "C"):
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["D"].width = 80
    ws["B4"].alignment = Alignment(wrap_text=True)

    # Static fallback so the file is useful even before xlwings is wired up.
    ws["A15"] = "Static snapshot (computed now — install xlwings for live cells):"
    ws["A15"].font = Font(italic=True, color="555555")
    fills = {"GO": "C6EFCE", "CONDITIONAL": "FFEB9C", "NO-GO": "FFC7CE"}
    ws.append([])  # spacer not strictly at row, but keep simple
    base = 16
    ws.cell(base, 1, "Deal").font = hf
    ws.cell(base, 1).fill = hdr
    ws.cell(base, 2, "Decision").font = hf
    ws.cell(base, 2).fill = hdr
    ws.cell(base, 3, "Score").font = hf
    ws.cell(base, 3).fill = hdr
    ws.cell(base, 4, "Why").font = hf
    ws.cell(base, 4).fill = hdr
    for j, (name, v) in enumerate(sorted(lib.items(), key=lambda kv: -kv[1].score)):
        r = base + 1 + j
        ws.cell(r, 1, v.name)
        c = ws.cell(r, 2, v.decision)
        c.fill = PatternFill("solid", fgColor=fills.get(v.decision, "FFFFFF"))
        ws.cell(r, 3, round(v.score, 1))
        ws.cell(r, 4, v.rationale)

    # --- Setup tab ----------------------------------------------------------
    ws2 = wb.create_sheet("Setup")
    steps = [
        ["How to make the cells generate LIVE answers"],
        [""],
        ["1. Install:  pip install \"xlwings>=0.30\"  and  pip install -e .  (this repo)"],
        ["2. In a terminal:  xlwings addin install"],
        ["3. Open this workbook in Excel. On the xlwings ribbon tab:"],
        ["   - set 'UDF Modules' to:  deal_agent.gonogo.excel"],
        ["   - click 'Import Functions'."],
        ["4. Tell it where your deals live (one of):"],
        ["   - set env var GONOGO_DEAL_FILES to a comma-separated list of files, OR"],
        ["   - drop your deal .xlsx/.csv files into a ./deals folder."],
        ["5. Now type any question in the Ask tab, e.g.  =GONOGO_ASK(\"is TX a go?\")"],
        [""],
        ["No Excel / prefer the terminal?  Ask from the command line:"],
        ["   python -m deal_agent.gonogo ask \"is Hamburg a go or no go?\""],
        [""],
        ["Hamburg is grounded in real data; Lewiston is a placeholder — edit"],
        ["deal_agent/gonogo/benchmarks.py to use your real Lewiston numbers."],
        [""],
        ["Analytical decision support — not legal, financial, or investment advice."],
    ]
    for row in steps:
        ws2.append(row)
    ws2.cell(1, 1).font = Font(bold=True, size=14)
    ws2.column_dimensions["A"].width = 95

    wb.save(out_path)
    return out_path
