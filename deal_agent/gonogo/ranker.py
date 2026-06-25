"""Rank deals by Go/No-Go score and write the ranking workbook.

`rank_deals` scores every deal and sorts GO > CONDITIONAL > NO-GO, then by
score. `write_ranking_workbook` renders it to a color-coded .xlsx with a
ranked summary tab and a per-metric detail tab — the "go no go ranking list".
"""

from __future__ import annotations

from .benchmarks import BENCHMARKS
from .engine import Verdict, score_deal
from .metrics import DealMetrics, extract_deals

_ORDER = {"GO": 0, "CONDITIONAL": 1, "NO-GO": 2}


def rank_deals(deals: list[DealMetrics]) -> list[Verdict]:
    """Score and rank. Returns verdicts best-first."""
    verdicts = [score_deal(d) for d in deals]
    verdicts.sort(key=lambda v: (_ORDER.get(v.decision, 3), -v.score))
    return verdicts


def rank_files(paths: list[str]) -> list[Verdict]:
    """Extract deals from every path, then rank them together."""
    deals: list[DealMetrics] = []
    for p in paths:
        deals.extend(extract_deals(p))
    return rank_deals(deals)


def _benchmark_verdicts() -> list[Verdict]:
    """The reference deals scored through the same engine (sanity anchor rows)."""
    rows = []
    for b in BENCHMARKS.values():
        m = DealMetrics(
            name=f"★ {b.name} (benchmark)",
            source="benchmark",
            yield_on_cost=b.yield_on_cost,
            exit_cap=b.exit_cap,
            dev_spread_bps=b.dev_spread_bps,
            levered_irr=b.levered_irr,
            unlevered_irr=b.unlevered_irr,
            moic=b.moic,
            lease_up_months=b.lease_up_months,
            noi=b.noi,
            total_cost=b.total_cost,
        )
        rows.append(score_deal(m))
    return rows


def write_ranking_workbook(
    verdicts: list[Verdict],
    out_path: str,
    include_benchmarks: bool = True,
) -> str:
    """Write a color-coded Go/No-Go ranking workbook. Returns out_path."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {
        "GO": PatternFill("solid", fgColor="C6EFCE"),
        "CONDITIONAL": PatternFill("solid", fgColor="FFEB9C"),
        "NO-GO": PatternFill("solid", fgColor="FFC7CE"),
    }
    fonts = {
        "GO": Font(color="006100", bold=True),
        "CONDITIONAL": Font(color="9C6500", bold=True),
        "NO-GO": Font(color="9C0006", bold=True),
    }
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    wb = Workbook()

    # --- Sheet 1: Ranking ---------------------------------------------------
    ws = wb.active
    ws.title = "Go-No-Go Ranking"
    bench_rows = _benchmark_verdicts() if include_benchmarks else []
    all_rows = verdicts + bench_rows
    cols = list(all_rows[0].as_row().keys()) if all_rows else []
    headers = ["Rank", *cols]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(1, c).fill = header_fill
        ws.cell(1, c).font = header_font
        ws.cell(1, c).alignment = Alignment(horizontal="center", wrap_text=True)

    rank = 0
    for v in all_rows:
        is_bench = v.name.startswith("★")
        if not is_bench:
            rank += 1
        row = v.as_row()
        ws.append(["BENCH" if is_bench else rank, *row.values()])
        r = ws.max_row
        dec = v.decision
        ws.cell(r, headers.index("Decision") + 1).fill = fills.get(dec)
        ws.cell(r, headers.index("Decision") + 1).font = fonts.get(dec)
        if is_bench:
            for c in range(1, len(headers) + 1):
                ws.cell(r, c).font = Font(italic=True, color="555555")

    _autosize(ws, get_column_letter, cap=48)
    ws.freeze_panes = "A2"

    # --- Sheet 2: Metric detail --------------------------------------------
    ws2 = wb.create_sheet("Score Detail")
    ws2.append(["Deal", "Decision", "Score", "Metric", "Value", "Band", "Points", "Knockout"])
    for c in range(1, 9):
        ws2.cell(1, c).fill = header_fill
        ws2.cell(1, c).font = header_font
    for v in verdicts:
        for b in v.breakdown:
            ws2.append([
                v.name, v.decision, round(v.score, 1),
                b.metric.replace("_", " ").title(),
                round(b.value, 4), b.band, round(b.points, 0),
                "YES" if b.knockout else "",
            ])
            ws2.cell(ws2.max_row, 6).fill = fills.get(b.band)
    _autosize(ws2, get_column_letter, cap=28)
    ws2.freeze_panes = "A2"

    # --- Sheet 3: Method / benchmarks --------------------------------------
    ws3 = wb.create_sheet("Method & Benchmarks")
    notes = [
        ["Go / No-Go Ranking — Method"],
        [""],
        ["Every deal is scored 0-100 against the team's proven reference deals:"],
        *[[f"  {b.name}: YoC {b.yield_on_cost:.1%}, exit cap {b.exit_cap:.2%}, "
           f"spread {b.dev_spread_bps:.0f} bps, levered IRR {b.levered_irr:.1%}, "
           f"MOIC {b.moic:.2f}x" + ("" if b.grounded else "  [PLACEHOLDER — edit benchmarks.py]")]
          for b in BENCHMARKS.values()],
        [""],
        ["Decision bands:"],
        ["  GO          score >= 70 and no metric below its hard floor"],
        ["  CONDITIONAL 55-70, clears every floor but trails the benchmarks"],
        ["  NO-GO       < 55, or any metric at/below its hard knockout floor"],
        [""],
        ["Weights: Dev Spread 26%, Yield on Cost 20%, Levered IRR 20%,"],
        ["         MOIC 12%, Market Score 12%, Lease-Up 6%, Exit Cap 4%."],
        [""],
        ["Analytical decision support — not legal, financial, or investment advice."],
    ]
    for row in notes:
        ws3.append(row)
    ws3.cell(1, 1).font = Font(bold=True, size=14)
    ws3.column_dimensions["A"].width = 90

    wb.save(out_path)
    return out_path


def _autosize(ws, get_column_letter, cap: int = 40) -> None:
    for col_cells in ws.columns:
        width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(width + 2, cap)
