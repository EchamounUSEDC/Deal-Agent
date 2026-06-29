"""Build the consolidated Co-Gen CUP presentation on top of the IC deck template.

- Inherits the template theme / master / layouts (navy / cream / red, Calibri).
- First paragraph of the memo becomes the first content slide.
- Each memo bullet point gets its own slide.
- Images are placed on the slides whose bullet they illustrate.
- An internal-comparison agent slide is included and the agent .py is embedded
  inside the .pptx package.
"""
import os
import shutil
import zipfile
import sys

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))      # the presentation/ folder
REPO = os.path.dirname(HERE)
IMG = f"{HERE}/assets"
OUTDIR = HERE
ASSETS = f"{HERE}/assets"
AGENT_SRC = f"{REPO}/deal_agent/agents/cogen_comparison_agent.py"
TEMPLATE = f"{HERE}/template.pptx"
OUT = f"{OUTDIR}/CoGen_CUP_Presentation.pptx"

os.makedirs(ASSETS, exist_ok=True)

NAVY = RGBColor(0x0F, 0x11, 0x3C)
RED = RGBColor(0xAA, 0x22, 0x26)
RED2 = RGBColor(0xC6, 0x00, 0x00)
CREAM = RGBColor(0xF3, 0xF2, 0xDC)
SLATE = RGBColor(0x80, 0x8D, 0xA0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DK = RGBColor(0x29, 0x29, 0x34)

prs = Presentation(TEMPLATE)
SW, SH = prs.slide_width, prs.slide_height

# Remove the 7 example slides that ship with the template (keep layouts/master).
# Drop each slide's relationship so the orphaned part is not re-serialized.
from pptx.oxml.ns import qn  # noqa: E402

xml_slides = prs.slides._sldIdLst
for sid in list(xml_slides):
    rId = sid.get(qn("r:id"))
    if rId:
        prs.part.drop_rel(rId)
    xml_slides.remove(sid)

layouts = {l.name: l for l in prs.slide_layouts}


def layout(name):
    return layouts[name]


def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    return None


def set_text(tf, text, size=None, color=None, bold=None, align=None):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    r = p.runs[0]
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if bold is not None:
        r.font.bold = bold
    if align:
        p.alignment = align
    return p


def no_bullet(p):
    """Strip any inherited bullet glyph from a paragraph (for prose, not lists)."""
    from pptx.oxml.ns import qn as _qn
    from pptx.oxml import parse_xml
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
        for e in pPr.findall(_qn(tag)):
            pPr.remove(e)
    # buNone must precede any defRPr in the schema order
    bun = parse_xml('<a:buNone xmlns:a="http://schemas.openxmlformats.org/'
                    'drawingml/2006/main"/>')
    defRPr = pPr.find(_qn("a:defRPr"))
    if defRPr is not None:
        defRPr.addprevious(bun)
    else:
        pPr.append(bun)


def add_bullets(tf, items, size=16, color=DK, space_after=8):
    """items: list of (text, level) or text. First fills paragraph[0]."""
    tf.word_wrap = True
    first = True
    for it in items:
        if isinstance(it, tuple):
            txt, lvl = it
        else:
            txt, lvl = it, 0
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = txt
        p.level = lvl
        p.space_after = Pt(space_after)
        for r in p.runs:
            r.font.size = Pt(size - (1 if lvl else 0))
            r.font.color.rgb = color


def style_footer(slide, footer_text="USGPD  ·  Co-Generation Central Utility Plant"):
    f = ph(slide, 11)
    if f is not None:
        set_text(f.text_frame, footer_text, size=9, color=SLATE)
    d = ph(slide, 10)
    if d is not None:
        set_text(d.text_frame, "June 29, 2026", size=9, color=SLATE)


def fit_box(img_path, box_w, box_h):
    """Return (w, h) EMU fitting img into box preserving aspect ratio."""
    with Image.open(img_path) as im:
        iw, ih = im.size
    ar = iw / ih
    bw, bh = box_w, box_h
    if bw / bh > ar:
        h = bh
        w = int(bh * ar)
    else:
        w = bw
        h = int(bw / ar)
    return w, h


def add_picture_fit(slide, img_path, left, top, box_w, box_h, frame=True):
    w, h = fit_box(img_path, box_w, box_h)
    l = left + (box_w - w) // 2
    t = top + (box_h - h) // 2
    pic = slide.shapes.add_picture(img_path, l, t, w, h)
    if frame:
        pic.line.color.rgb = SLATE
        pic.line.width = Pt(0.75)
    return pic


CONTENT_LAYOUT = "1_Title and Content"


def content_slide(title, main_point, bullets=None, body=None, body_size=16):
    """Standard text slide using the template's content layout."""
    s = prs.slides.add_slide(layout(CONTENT_LAYOUT))
    set_text(ph(s, 0).text_frame, title)                 # Title (inherits style)
    mp = ph(s, 13)
    if mp is not None:
        set_text(mp.text_frame, main_point, color=RED)
    c = ph(s, 1)
    if bullets:
        add_bullets(c.text_frame, bullets, size=body_size)
    elif body:
        # Prose paragraph: no bullet glyph, generous size, anchored near the top.
        c.top = Inches(1.7)
        c.height = Inches(4.6)
        tf = c.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = set_text(tf, body, size=body_size, color=DK)
        p.line_spacing = 1.15
        p.space_after = Pt(10)
        no_bullet(p)
    style_footer(s)
    return s


def image_slide(title, main_point, bullets, img_path, caption=None, body_size=15):
    """Text on the left half, image on the right half."""
    s = prs.slides.add_slide(layout(CONTENT_LAYOUT))
    set_text(ph(s, 0).text_frame, title)
    mp = ph(s, 13)
    if mp is not None:
        set_text(mp.text_frame, main_point, color=RED)
    # shrink content placeholder to the left column
    c = ph(s, 1)
    c.left = Inches(0.6)
    c.top = Inches(1.7)
    c.width = Inches(6.0)
    c.height = Inches(4.9)
    add_bullets(c.text_frame, bullets, size=body_size)
    # image on the right column
    img_left = Inches(6.9)
    img_top = Inches(1.7)
    img_box_w = Inches(6.0)
    img_box_h = Inches(4.6)
    add_picture_fit(s, img_path, img_left, img_top, img_box_w, img_box_h)
    if caption:
        tb = s.shapes.add_textbox(img_left, Inches(6.35), img_box_w, Inches(0.4))
        set_text(tb.text_frame, caption, size=9, color=SLATE, align=PP_ALIGN.CENTER)
    style_footer(s)
    return s


def divider_slide(title, main_point, items):
    s = prs.slides.add_slide(layout(CONTENT_LAYOUT))
    set_text(ph(s, 0).text_frame, title)
    mp = ph(s, 13)
    if mp is not None:
        set_text(mp.text_frame, main_point, color=RED)
    c = ph(s, 1)
    add_bullets(c.text_frame, items, size=18, space_after=10)
    style_footer(s)
    return s


# ============================================================ SLIDES ===========

# --- 1. TITLE ----------------------------------------------------------------
t = prs.slides.add_slide(layout("1 Title Slide"))
if ph(t, 10) is not None:
    set_text(ph(t, 10).text_frame, "June 29, 2026")
if ph(t, 11) is not None:
    set_text(ph(t, 11).text_frame, "Central Utility Plant (CUP) — Co-Generation Project")
if ph(t, 12) is not None:
    set_text(ph(t, 12).text_frame, "Mammoth Hill  ·  Central City, Colorado")
notes = t.notes_slide.notes_text_frame
notes.text = ("Consolidated Co-Gen CUP presentation. Source memo: 'Goals and Timeline — "
              "USGPD as Developer and Operator of the Central Utility Plant', prepared by "
              "Bill Hibbard, 6/5/26, plus the Aaron Equipment Caterpillar CG170-16 quote.")

# --- 2. PROJECT OVERVIEW (first paragraph of the memo) -----------------------
overview = ("The Co-Gen Special Project Group as well as PDG have considered the recent "
            "delivery options as well as progress completed to date. Development of "
            "construction documents for the CUP is critical to completing the application "
            "and approval for the Xcel gas use application. We provide the following points "
            "for consideration by ownership to continue on our current path to include the "
            "CUP as part of the overall delivery of the Casino project.")
content_slide("Project Overview",
              "USGPD as Developer & Operator of the Central Utility Plant (CUP)",
              body=overview, body_size=20)

# --- 3. DIVIDER: ADVANTAGES --------------------------------------------------
divider_slide("Advantages of Developing, Owning & Operating the CUP",
              "Seven reasons to keep the CUP inside the Casino project delivery",
              ["Primary Focus", "Application Timeline", "CUP Location",
               "Alternate Delivery Sources", "Utility Cost Control",
               "Long-Term Economic Payback", "Central City Mini-District Constraints"])

# --- 4. Primary Focus --------------------------------------------------------
content_slide("Primary Focus", "The CUP is essential to the overall project delivery",
              bullets=[
                  "Our primary focus is the construction and operation of the Casino.",
                  "The CUP is an essential part of the overall project delivery and "
                  "long-term operation.",
                  "Being one user in a larger district would transfer timely delivery and "
                  "key decision-making to a separate entity.",
              ])

# --- 5. Application Timeline -------------------------------------------------
content_slide("Application Timeline",
              "IMEG selected (May) to deliver the Feasibility Study; Xcel approval path",
              bullets=[
                  "A Co-Gen consultant was identified in early April for a Feasibility "
                  "Study; the bid process completed in May and IMEG was selected.",
                  ("IMEG Phase I — ROM & Schedule:  6/5/26 → 6/19/26", 1),
                  ("IMEG Phase II — Equipment Selection, P&ID, Floor Plan, Economic "
                   "Payback Analysis:  6/19/26 → 8/3/26", 1),
                  ("Energy Modeling (IMEG or Group 14):  7/15/26 → 9/15/26", 1),
                  ("Prepare Construction Documents for the CUP:  8/3/26 → 10/30/26", 1),
                  ("Submit CDs & Application to Xcel for approval:  11/2/26", 1),
              ], body_size=15)

# --- 6. CUP Location  [IMG site map] -----------------------------------------
image_slide("CUP Location",
            "Mammoth Hill — proximity is key to reliability and cost",
            bullets=[
                "Proximity of the CUP to the Casino site maximizes reliability and "
                "minimizes risk.",
                "A steam / condensate pipe bridge delivers superheated steam to the "
                "Casino Mechanical Room and returns condensate to the CUP.",
                "Minimizing pipe-bridge length (pump transfer stations, traps, etc.) "
                "reduces maintenance, shutdown risk, service interruption and cost.",
                "Mammoth Hill is the primary location for the CUP site selection process.",
            ],
            img_path=f"{IMG}/site_map.png",
            caption="Schematic site context — Mammoth Hill CUP relative to the Casino")

# --- 7. Alternate Delivery Sources ------------------------------------------
content_slide("Alternate Delivery Sources",
              "Grid-only delivery adds maintenance, risk and a ~5-year Xcel timeline",
              bullets=[
                  "Special Projects met with Xcel on options for delivering electricity "
                  "to the Casino project.",
                  "Current connected load ≈ 8.5 MW;  regular operating load ≈ 5–6.5 MW.",
                  "Xcel can provide 2.5 MW without improvements to the current grid "
                  "infrastructure.",
                  "Additional power could be run from Idaho Springs (underground or "
                  "overhead) — adding maintenance, reliability concerns and outage risk.",
                  "Xcel reported installation would require 5 years from application.",
              ], body_size=15)

# --- 8. Utility Cost Control -------------------------------------------------
content_slide("Utility Cost Control",
              "Internal cost management vs. payments to a District in perpetuity",
              bullets=[
                  "Cost control of the generated electricity will be managed by USGPD.",
                  "Market rate for electricity ≈ $0.07/kWh (region-dependent, excludes "
                  "transmission);  natural gas ≈ $4.50/DTH.",
                  "The CUP cost per kWh will be a cost-of-operation figure set by the "
                  "economic payback analysis.",
                  "Internal management of this cost — versus payments to a District in "
                  "perpetuity — is an important factor in the decision.",
              ])

# --- 9. Long-Term Economic Payback  [IMG combined cycle] ---------------------
image_slide("Long-Term Economic Payback",
            "‘Free heating & cooling’ from turbine stack discharge",
            bullets=[
                "The Casino will be the primary, if not sole, user of the thermal energy "
                "(superheated steam) generated by the CUP.",
                "This provides 'free heating and cooling' for the Resort by optimizing "
                "the turbine generators' high-temperature stack discharge.",
                "Overall cost of energy to operate the Casino will be determined in the "
                "IMEG Feasibility Study and Economic Payback Analysis.",
            ],
            img_path=f"{IMG}/combined_cycle.png",
            caption="Combined-cycle co-generation: one fuel, electricity + thermal")

# --- 10. Central City Mini-District Constraints ------------------------------
content_slide("Central City Mini-District Constraints",
              "A third-party district adds delay, perpetual cost and distance risk",
              bullets=[
                  "Central City would most likely hire a third-party Developer/Operator.",
                  "This would delay the construction-document and gas-application process.",
                  "It would lead to negotiation of utility costs in perpetuity.",
                  "A district plant farther from the Casino adds substantial cost (pipe-"
                  "bridge length) and raises reliability and disruption risk.",
              ])

# --- 11. DIVIDER: DISADVANTAGE -----------------------------------------------
divider_slide("Disadvantage to Consider",
              "The principal trade-off of owning and operating the CUP",
              ["Time Value of Money"])

# --- 12. Time Value of Money -------------------------------------------------
content_slide("Time Value of Money",
              "A $24M–$30M CUP carries capital and interest before full operation",
              bullets=[
                  "The CUP is estimated to cost between $24M and $30M.",
                  "An independent owner/operator could reduce the total funds USGPD raises "
                  "through syndication and the interest due on funds prior to full "
                  "operation and economic return of the CUP.",
                  "The economic payback analysis (IMEG) is part of the overall "
                  "feasibility analysis.",
              ])

# --- 13. Proposed Equipment  [IMG Caterpillar] -------------------------------
image_slide("Proposed Equipment — Caterpillar CG170-16",
            "Unused 1.5 MW natural-gas cogeneration generator set — $599,000",
            bullets=[
                "Unused Caterpillar Cogeneration Natural Gas Generator System, 1.5 MW.",
                "CG170-16 engine rated 1,556 kW, 4160 V, 3-phase, 60 Hz at 0.8 PF.",
                "Pritchard Brown enclosure, Cannon boiler, Boulden radiator, WHR system, "
                "Stueler SCR system (outside enclosure), Caterpillar ISO aux cabinet.",
                "Built 2017;  80 hours of engineering support to site.  U.S.A. only.",
                ("Item price:  $599,000.00 USD", 0),
            ],
            img_path=f"{IMG}/49ded102_image1.jpeg",
            caption="Aaron Equipment — Stock #50477001 (Caterpillar CG170-16)",
            body_size=14)

# --- 14. Emissions Control & Exhaust Dilution  [IMG plume] -------------------
image_slide("Emissions Control & Exhaust Dilution",
            "SCR-controlled stack discharge — plume dilution with downwind distance",
            bullets=[
                "The Stueler SCR system controls exhaust emissions from the gas engine.",
                "Stack discharge is the high-temperature heat source captured for the "
                "Casino's thermal load.",
                "Plume-concentration modeling (e.g., Greenheck CAPS-style analysis) "
                "confirms concentrations fall rapidly with downwind distance and height.",
                "Point of interest at 120 ft downwind / 10 ft high ≈ 1,392 PPM — used to "
                "site the stack and intakes.",
            ],
            img_path=f"{IMG}/plume.png",
            caption="Illustrative exhaust-plume dilution (concentration bands)",
            body_size=14)

# --- 15. Project Update & Status ---------------------------------------------
content_slide("Project Update & Status to Date",
              "Xcel confirms gas and 2.5 MW available without infrastructure upgrades",
              bullets=[
                  "USGPD and the design team (Martin & Martin, OS, RTM) have been in "
                  "discussions with Xcel to apply for gas and electric service.",
                  "Xcel has communicated that 9 million cubic feet of gas and 2.5 MW of "
                  "electricity are available for the Casino project (see Meeting Minutes).",
                  "These quantities may be provided without improvements to area utility "
                  "infrastructure.",
              ])

# --- 16. RFP & Consultant Selection ------------------------------------------
content_slide("RFP & Consultant Selection",
              "Fast-track strategy: complete CDs required before Xcel approval",
              bullets=[
                  "On April 7th, USGPD issued an RFP to four Co-Gen specialist consultants "
                  "for a Feasibility Study and Economic Payback Analysis.",
                  "Proposals were received April 22nd.",
                  "NOTE — Applications for electric and gas service have been submitted to "
                  "Xcel. Xcel requires a complete set of Construction Documents before "
                  "approving service and quantities — driving a 'fast-track' strategy.",
              ])

# --- 17. Recommendation ------------------------------------------------------
rec = prs.slides.add_slide(layout(CONTENT_LAYOUT))
set_text(ph(rec, 0).text_frame, "Recommendation")
set_text(ph(rec, 13).text_frame,
         "Include the CUP as part of the overall Casino project", color=RED)
add_bullets(ph(rec, 1).text_frame, [
    "Central City and independent developers have expressed interest in owning and "
    "operating a Utility Mini-District to serve the Casino and future development.",
    "Due to the factors above — control, timeline, reliability, cost control and "
    "long-term thermal payback — the current recommendation is to develop, own and "
    "operate the CUP within the Casino project.",
    ("Prepared by Bill Hibbard  ·  USGPD Special Projects  ·  6/5/26", 0),
], size=17)
style_footer(rec)

# --- 18. INTERNAL COMPARISON AGENT -------------------------------------------
import importlib.util
spec = importlib.util.spec_from_file_location("cogen_comparison_agent", AGENT_SRC)
_agent_mod = importlib.util.module_from_spec(spec)
sys.modules["cogen_comparison_agent"] = _agent_mod
spec.loader.exec_module(_agent_mod)
data = _agent_mod.run_internal_comparison()
agent_slide = prs.slides.add_slide(layout("Title Only"))
set_text(ph(agent_slide, 0).text_frame, "Internal Comparison Agent")

# intro line
tb = agent_slide.shapes.add_textbox(Inches(0.5), Inches(0.95), Inches(12.3), Inches(0.5))
set_text(tb.text_frame,
         "A dependency-free agent (embedded in this file) scores the three CUP "
         "delivery options against weighted criteria.", size=13, color=DK)

# build the comparison table
crit = data["criteria"]
res = data["results"]
short = {"self_develop": "Self-Develop\n(USGPD)", "mini_district": "Mini-District\n(3rd party)",
         "grid_only": "Grid-Only\n(Xcel)"}
order = [r["key"] for r in res]  # already sorted best-first
rows = len(crit) + 2  # header + criteria + weighted total
cols = 2 + len(order)
tbl_left, tbl_top, tbl_w, tbl_h = Inches(0.5), Inches(1.6), Inches(12.3), Inches(4.6)
gtbl = agent_slide.shapes.add_table(rows, cols, tbl_left, tbl_top, tbl_w, tbl_h).table
gtbl.columns[0].width = Inches(5.0)
gtbl.columns[1].width = Inches(1.3)
for i in range(len(order)):
    gtbl.columns[2 + i].width = Inches(2.0)

def cell(r, c, text, size=12, color=DK, bold=False, fill=None, align=PP_ALIGN.CENTER):
    cl = gtbl.cell(r, c)
    cl.margin_top = Pt(2); cl.margin_bottom = Pt(2)
    cl.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is not None:
        cl.fill.solid(); cl.fill.fore_color.rgb = fill
    tf = cl.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    for rn in p.runs:
        rn.font.size = Pt(size); rn.font.color.rgb = color; rn.font.bold = bold

# header
cell(0, 0, "Criterion", bold=True, color=WHITE, fill=NAVY, align=PP_ALIGN.LEFT)
cell(0, 1, "Weight", bold=True, color=WHITE, fill=NAVY)
for i, k in enumerate(order):
    cell(0, 2 + i, short[k], bold=True, color=WHITE, fill=NAVY)
# criteria rows
for ri, c in enumerate(crit, start=1):
    band = CREAM if ri % 2 else WHITE
    cell(ri, 0, c["name"], color=DK, fill=band, align=PP_ALIGN.LEFT)
    cell(ri, 1, f'{c["weight"]}%', color=DK, fill=band)
    for i, k in enumerate(order):
        cell(ri, 2 + i, str(c["scores"][k]), color=DK, fill=band)
# weighted total row
tr = len(crit) + 1
cell(tr, 0, "WEIGHTED SCORE  (1–5, higher = better)", bold=True, color=WHITE, fill=RED,
     align=PP_ALIGN.LEFT)
cell(tr, 1, "100%", bold=True, color=WHITE, fill=RED)
for i, k in enumerate(order):
    sc = next(r["weighted_score"] for r in res if r["key"] == k)
    cell(tr, 2 + i, f"{sc:.2f}", bold=True, color=WHITE, fill=RED)

# recommendation strip
rb = agent_slide.shapes.add_textbox(Inches(0.5), Inches(6.35), Inches(12.3), Inches(0.6))
set_text(rb.text_frame, "▶ " + data["recommendation"], size=13, color=NAVY, bold=True)

agent_slide.notes_slide.notes_text_frame.text = (
    "This deck embeds a runnable agent: deal_agent/agents/cogen_comparison_agent.py is "
    "also stored inside this .pptx package under /agent/. Run it for the live matrix:\n"
    "    python cogen_comparison_agent.py            # formatted table + recommendation\n"
    "    python cogen_comparison_agent.py --json     # machine-readable scoring matrix\n"
    "Criteria weights: Control 20, Timeline 20, Reliability 20, Cost Control 15, "
    "Capital 15, Thermal Payback 10. Scores are 1-5 from USGPD's goal perspective.")

# ---------------------------------------------------------------------------
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides._sldIdLst))

# copy assets used into the deliverable folder (skip if already co-located)
for fn in ["site_map.png", "combined_cycle.png", "plume.png", "49ded102_image1.jpeg"]:
    src, dst = f"{IMG}/{fn}", f"{ASSETS}/{fn}"
    if os.path.abspath(src) != os.path.abspath(dst):
        shutil.copy(src, dst)

# embed the agent .py physically inside the .pptx package.
# Rewrite the archive so [Content_Types].xml declares a Default for ".py"
# (keeps the OPC package valid / PowerPoint-openable with an extra part inside).
EMBED_PATH = "agent/cogen_comparison_agent.py"
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data_bytes = zin.read(item.filename)
        if item.filename == "[Content_Types].xml":
            ct = data_bytes.decode()
            if 'Extension="py"' not in ct:
                ct = ct.replace(
                    "</Types>",
                    '<Default Extension="py" ContentType="text/x-python"/></Types>')
            data_bytes = ct.encode()
        zout.writestr(item, data_bytes)
    with open(AGENT_SRC, "rb") as f:
        zout.writestr(EMBED_PATH, f.read())
os.replace(tmp, OUT)
print("embedded agent at", EMBED_PATH)
